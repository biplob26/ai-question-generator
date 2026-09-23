import streamlit as st
import google.generativeai as genai
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

# Streamlit পেজ সেটিং
st.set_page_config(page_title="AI Question Paper Generator", layout="wide")

st.title("📝 AI প্রশ্নপত্র ও উত্তরপত্র জেনারেটর")
st.write("প্রাথমিক ও মাধ্যমিক স্তরের জন্য প্রিন্ট-রেডি প্রশ্নপত্র এবং উত্তরপত্র তৈরি করুন।")

# MS Word ফাইল তৈরির ফাংশন
def create_word_docx(content, title_info):
    doc = Document()
    
    # পেজ মার্জিন সেটিং (০.৭৫ ইঞ্চি)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # হেডার / শিরোনাম
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_head.add_run(f"{title_info['institution']}\n")
    run_inst.bold = True
    run_inst.font.size = Pt(16)
    
    run_exam = p_head.add_run(f"{title_info['exam']}\n")
    run_exam.bold = True
    run_exam.font.size = Pt(13)
    
    run_meta = p_head.add_run(f"শ্রেণি: {title_info['class_name']} | বিষয়: {title_info['subject']}\n")
    run_meta.font.size = Pt(11)
    
    p_info = doc.add_paragraph()
    p_info.add_run(f"সময়: {title_info['time']}                                                      মোট নম্বর: {title_info['marks']}")
    p_info.runs[0].font.size = Pt(11)
    doc.add_paragraph("-" * 55)

    # মূল কন্টেন্ট যোগ করা
    for line in content.split("\n"):
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)
        
    # মেমোরি থেকে বাইট ফাইল আকারে রিটার্ন
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io

# সাইডবারে এপিআই কি ইনপুট
with st.sidebar:
    st.header("⚙️ সেটআপ")
    api_key = st.text_input("আপনার Google Gemini API Key দিন:", type="password")
    st.markdown("[ফ্রি API Key পেতে এখানে ক্লিক করুন](https://aistudio.google.com/)")

# ইনপুট ফর্ম
with st.form("question_form"):
    col1, col2 = st.columns(2)
    
    with col1:
        institution = st.text_input("প্রতিষ্ঠানের নাম", "আইডিয়াল স্কুল অ্যান্ড কলেজ")
        exam_name = st.text_input("পরীক্ষার নাম", "বার্ষিক পরীক্ষা - ২০২৬")
        level_category = st.selectbox("শিক্ষার স্তর নির্বাচন করুন", ["প্রাথমিক (১ম - ৫ম শ্রেণি)", "মাধ্যমিক (৬ষ্ঠ - ১০ম শ্রেণি)"])
        class_name = st.text_input("শ্রেণি", "৪র্থ" if "প্রাথমিক" in level_category else "৮ম")
        subject = st.text_input("বিষয়", "গণিত" if "প্রাথমিক" in level_category else "বিজ্ঞান")
        
    with col2:
        full_marks = st.number_input("মোট নম্বর", value=50, step=5)
        time_limit = st.text_input("সময়", "১ ঘণ্টা ৩০ মিনিট")
        chapters = st.text_area("অধ্যায়/টপিকসমূহ", "অধ্যায় ১: বড় সংখ্যা ও স্থানীয় মান\nঅধ্যায় ২: যোগ, বিয়োগ, গুণ ও ভাগ")
        include_answers = st.checkbox("প্রশ্নপত্রের নিচে উত্তরপত্র (Answer Key) যুক্ত করুন", value=True)

    st.subheader("প্রশ্নপত্রের ধরন")
    if "প্রাথমিক" in level_category:
        col3, col4, col5 = st.columns(3)
        with col3:
            short_q_count = st.number_input("সংক্ষিপ্ত প্রশ্ন", value=5, min_value=0)
        with col4:
            fill_blank_count = st.number_input("শূন্যস্থান পূরণ", value=5, min_value=0)
        with col5:
            broad_q_count = st.number_input("কাঠামোগত/রচনামূলক প্রশ্ন", value=4, min_value=0)
        mcq_count = 0
        cq_count = 0
    else:
        col3, col4 = st.columns(2)
        with col3:
            mcq_count = st.number_input("বহুনির্বাচনী প্রশ্ন (MCQ)", value=10, min_value=0)
        with col4:
            cq_count = st.number_input("সৃজনশীল প্রশ্ন (CQ)", value=3, min_value=0)
        short_q_count = 0
        fill_blank_count = 0
        broad_q_count = 0

    submit_button = st.form_submit_button("🚀 প্রশ্ন ও উত্তরপত্র তৈরি করুন")

# এআই প্রম্পট ও জেনারেশন
if submit_button:
    if not api_key:
        st.error("❌ অনুগ্রহ করে সাইডবারে আপনার Gemini API Key দিন।")
    else:
        try:
            genai.configure(api_key=api_key)
            model = genai.GenerativeModel('gemini-2.5-flash')

            # স্তর অনুযায়ী প্রম্পট কাস্টমাইজেশন
            if "প্রাথমিক" in level_category:
                question_structure = f"""
                ১. সংক্ষিপ্ত উত্তর প্রশ্ন: {short_q_count} টি।
                ২. শূন্যস্থান পূরণ: {fill_blank_count} টি।
                ৩. কাঠামোগত/যোগ্যতাভিত্তিক রচনামূলক প্রশ্ন: {broad_q_count} টি।
                """
            else:
                question_structure = f"""
                ১. বহুনির্বাচনী প্রশ্ন (MCQ): {mcq_count} টি (প্রতিটির ৪টি অপশন ক, খ, গ, ঘ থাকবে)।
                ২. সৃজনশীল প্রশ্ন (CQ): {cq_count} টি (ক, খ, গ, ঘ অংশসহ)।
                """

            answer_prompt_instruction = ""
            if include_answers:
                answer_prompt_instruction = """
                প্রশ্নপত্রের শেষে একটি পৃথক সেকশন তৈরি করো যার শিরোনাম হবে:
                '=========================================='
                'উত্তরপত্র (Answer Key / সমাধান)'
                '=========================================='
                এখানে প্রতিটি প্রশ্নের নম্বরসহ সঠিক উত্তর বা সংক্ষিপ্ত সমাধান সুন্দরভাবে উল্লেখ করো।
                """

            prompt = f"""
            তুমি বাংলাদেশের একজন অভিজ্ঞ শিক্ষক। নিচের তথ্য অনুযায়ী {level_category}-এর জন্য একটি সম্পূর্ণ এবং নির্ভুল প্রশ্নপত্র তৈরি করো।

            প্রতিষ্ঠানের নাম: {institution}
            পরীক্ষার নাম: {exam_name}
            শ্রেণি: {class_name} | বিষয়: {subject}
            সময়: {time_limit} | মোট নম্বর: {full_marks}
            অধ্যায়সমূহ: {chapters}

            প্রশ্নের কাঠামো:
            {question_structure}

            {answer_prompt_instruction}

            নির্দেশনা:
            ১. বাংলাদেশের জাতীয় শিক্ষাক্রমের নিয়ম ও মান বজায় রাখো।
            ২. ভাষা সাবলীল ও স্পষ্ট হতে হবে।
            ৩. কোনো অপ্রয়োজনীয় ভূমিকা বা কথা না লিখে সরাসরি প্রশ্নপত্র ও উত্তরপত্র জেনারেট করো।
            """

            with st.spinner("প্রশ্ন ও উত্তরপত্র তৈরি হচ্ছে..."):
                response = model.generate_content(prompt)
                generated_text = response.text
                
                st.success("✅ প্রশ্ন ও উত্তরপত্র সফলভাবে তৈরি হয়েছে!")
                st.markdown("---")
                st.markdown(generated_text)

                # Word ফাইল জেনারেট
                title_data = {
                    "institution": institution,
                    "exam": exam_name,
                    "class_name": class_name,
                    "subject": subject,
                    "time": time_limit,
                    "marks": full_marks
                }
                docx_file = create_word_docx(generated_text, title_data)

                # MS Word ডাউনলোড বাটন
                st.download_button(
                    label="📄 MS Word (.docx) ফাইল নামান",
                    data=docx_file,
                    file_name=f"Question_{class_name}_{subject}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

        except Exception as e:
            st.error(f"একটি সমস্যা হয়েছে: {e}")
